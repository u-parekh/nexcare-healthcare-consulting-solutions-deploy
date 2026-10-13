import io
from datetime import datetime, timedelta
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query, Request, Response
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from auth import COOKIE, check_lock, create_session, record_fail, require_admin, verify_password
from core import client_ip, db, limiter, now, pseudonymize, safe_regex
from notifications import deliver_contact
from validators import ASSESSMENT_STATUSES, CONTACT_STATUSES, LoginIn, RecordUpdate

router = APIRouter(prefix="/api/admin")

COLLECTIONS = {
    "contacts": {"coll": "contacts", "statuses": CONTACT_STATUSES, "search": ["name", "email", "phone", "message", "service"]},
    "assessments": {"coll": "quiz_submissions", "statuses": ASSESSMENT_STATUSES, "search": ["name", "email", "phone"]},
}
HIDDEN = {"_id": 0, "ip_hash": 0, "token_hash": 0, "msg_hash": 0}


async def audit(user: dict, action: str, target: str, detail: Optional[dict] = None):
    await db.audit_log.insert_one({"actor": user.get("username"), "action": action, "target": target,
                                   "detail": detail or {}, "created_at": now()})


def kind_cfg(kind: str) -> dict:
    if kind not in COLLECTIONS:
        raise HTTPException(status_code=404, detail="Unknown dataset")
    return COLLECTIONS[kind]


def build_query(cfg, search, status, service, date_from, date_to, recent_days):
    q = {}
    if search:
        rx = safe_regex(search)
        q["$or"] = [{f: rx} for f in cfg["search"]]
    if status:
        if status not in cfg["statuses"]:
            raise HTTPException(status_code=400, detail="Invalid status filter")
        q["status"] = status
    if service:
        q["service"] = service[:120]
    rng = {}
    try:
        if date_from:
            rng["$gte"] = datetime.fromisoformat(date_from)
        if date_to:
            rng["$lt"] = datetime.fromisoformat(date_to) + timedelta(days=1)
    except ValueError:
        raise HTTPException(status_code=400, detail="Invalid date")
    if recent_days:
        rng["$gte"] = now() - timedelta(days=recent_days)
    if rng:
        q["created_at"] = rng
    return q


def filters(search: str = Query("", max_length=100), status: str = Query("", max_length=20),
            service: str = Query("", max_length=120), date_from: str = Query("", max_length=10),
            date_to: str = Query("", max_length=10), recent_days: int = Query(0, ge=0, le=365),
            sort: str = Query("desc", pattern="^(asc|desc)$")):
    return dict(search=search, status=status, service=service, date_from=date_from,
                date_to=date_to, recent_days=recent_days, sort=sort)


@router.post("/login")
async def login(payload: LoginIn, request: Request, response: Response):
    ip_h = pseudonymize(client_ip(request))
    limiter.check(f"login:{ip_h}", 10, 300)
    identifier = f"{ip_h}:{payload.username.lower()}"
    await check_lock(identifier)
    user = await db.admin_users.find_one({"username": payload.username})
    if not user or not verify_password(payload.password, user["password_hash"]):
        await record_fail(identifier)
        await db.audit_log.insert_one({"actor": payload.username[:64], "action": "login_failed", "target": "admin", "detail": {}, "created_at": now()})
        raise HTTPException(status_code=401, detail="Invalid username or password")
    await db.login_attempts.delete_one({"identifier": identifier})
    await create_session(response, user, ip_h)
    await audit(user, "login", "admin")
    return {"username": user["username"], "role": user["role"]}


@router.post("/logout")
async def logout(response: Response, user: dict = Depends(require_admin)):
    await db.admin_sessions.delete_one({"jti": user["jti"]})
    response.delete_cookie(COOKIE, path="/api")
    await audit(user, "logout", "admin")
    return {"ok": True}


@router.get("/me")
async def me(user: dict = Depends(require_admin)):
    return {"username": user["username"], "role": user["role"]}


@router.get("/analytics")
async def analytics(_: dict = Depends(require_admin)):
    t = now()
    day0 = t.replace(hour=0, minute=0, second=0, microsecond=0)

    async def uniq(since):
        return len(await db.visits.distinct("visitor_id", {"created_at": {"$gte": since}}))

    series = await db.visits.aggregate([
        {"$match": {"created_at": {"$gte": day0 - timedelta(days=29)}}},
        {"$group": {"_id": {"d": {"$dateToString": {"format": "%Y-%m-%d", "date": "$created_at"}}, "v": "$visitor_id"}, "views": {"$sum": 1}}},
        {"$group": {"_id": "$_id.d", "visitors": {"$sum": 1}, "views": {"$sum": "$views"}}},
        {"$sort": {"_id": 1}},
    ]).to_list(40)
    by_day = {s["_id"]: s for s in series}
    days = [(day0 - timedelta(days=i)).strftime("%Y-%m-%d") for i in range(29, -1, -1)]

    async def group(coll, field, limit=8, match=None):
        pipe = ([{"$match": match}] if match else []) + [{"$group": {"_id": f"${field}", "count": {"$sum": 1}}}, {"$sort": {"count": -1}}, {"$limit": limit}]
        return [{"name": r["_id"] or "Unknown", "count": r["count"]} for r in await db[coll].aggregate(pipe).to_list(limit)]

    recent = await db.contacts.find({}, {**HIDDEN, "message": 0}).sort("created_at", -1).limit(5).to_list(5)
    geo_measured = await db.visits.count_documents({"geo_source": "ip_lookup"})
    return {
        "visitors": {
            "total_unique": len(await db.visits.distinct("visitor_id")),
            "total_pageviews": await db.visits.count_documents({}),
            "today": await uniq(day0), "week": await uniq(t - timedelta(days=7)), "month": await uniq(t - timedelta(days=30)),
        },
        "series": [{"date": d, "visitors": by_day.get(d, {}).get("visitors", 0), "views": by_day.get(d, {}).get("views", 0)} for d in days],
        "countries": await group("visits", "country"),
        "geo_measured_pageviews": geo_measured,
        "top_pages": await group("visits", "path", 10),
        "devices": await group("visits", "device", 5),
        "contacts": {"total": await db.contacts.count_documents({}), "by_status": await group("contacts", "status"),
                     "by_service": await group("contacts", "service", 12)},
        "assessments": {"total": await db.quiz_submissions.count_documents({}), "by_status": await group("quiz_submissions", "status")},
        "brochure_downloads": await db.brochure_downloads.count_documents({}),
        "recent_contacts": recent,
    }


@router.get("/audit")
async def audit_list(page: int = Query(1, ge=1), per_page: int = Query(20, ge=1, le=100), _: dict = Depends(require_admin)):
    total = await db.audit_log.count_documents({})
    items = await db.audit_log.find({}, {"_id": 0}).sort("created_at", -1).skip((page - 1) * per_page).limit(per_page).to_list(per_page)
    return {"items": items, "total": total, "page": page, "pages": max(1, -(-total // per_page))}


@router.get("/records/{kind}")
async def list_records(kind: str, page: int = Query(1, ge=1), per_page: int = Query(10, ge=1, le=100),
                       f: dict = Depends(filters), _: dict = Depends(require_admin)):
    cfg = kind_cfg(kind)
    sort = f.pop("sort")
    q = build_query(cfg, **f)
    total = await db[cfg["coll"]].count_documents(q)
    items = await db[cfg["coll"]].find(q, HIDDEN).sort("created_at", 1 if sort == "asc" else -1).skip((page - 1) * per_page).limit(per_page).to_list(per_page)
    return {"items": items, "total": total, "page": page, "per_page": per_page, "pages": max(1, -(-total // per_page))}


@router.get("/records/{kind}/{rid}")
async def get_record(kind: str, rid: str, _: dict = Depends(require_admin)):
    doc = await db[kind_cfg(kind)["coll"]].find_one({"id": rid[:64]}, HIDDEN)
    if not doc:
        raise HTTPException(status_code=404, detail="Not found")
    return doc


@router.patch("/records/{kind}/{rid}")
async def update_record(kind: str, rid: str, payload: RecordUpdate, user: dict = Depends(require_admin)):
    cfg = kind_cfg(kind)
    upd = {}
    if payload.status is not None:
        if payload.status not in cfg["statuses"]:
            raise HTTPException(status_code=400, detail="Invalid status")
        upd["status"] = payload.status
    if payload.notes is not None:
        upd["admin_notes"] = payload.notes.strip()
    if not upd:
        raise HTTPException(status_code=400, detail="Nothing to update")
    upd["updated_at"] = now()
    res = await db[cfg["coll"]].update_one({"id": rid[:64]}, {"$set": upd})
    if res.matched_count == 0:
        raise HTTPException(status_code=404, detail="Not found")
    await audit(user, f"update_{kind}", rid, {k: v for k, v in upd.items() if k != "updated_at"})
    return await db[cfg["coll"]].find_one({"id": rid[:64]}, HIDDEN)


@router.delete("/records/{kind}/{rid}")
async def delete_record(kind: str, rid: str, user: dict = Depends(require_admin)):
    res = await db[kind_cfg(kind)["coll"]].delete_one({"id": rid[:64]})
    if res.deleted_count == 0:
        raise HTTPException(status_code=404, detail="Not found")
    await audit(user, f"delete_{kind}", rid)
    return {"deleted": True}


@router.post("/records/contacts/{rid}/retry-delivery")
async def retry_delivery(rid: str, user: dict = Depends(require_admin)):
    c = await db.contacts.find_one({"id": rid[:64]}, HIDDEN)
    if not c:
        raise HTTPException(status_code=404, detail="Not found")
    delivery = await deliver_contact(c, c.get("delivery"))
    await db.contacts.update_one({"id": c["id"]}, {"$set": {"delivery": delivery, "updated_at": now()}})
    await audit(user, "retry_delivery", rid, {k: v["status"] for k, v in delivery.items()})
    return {"delivery": delivery}


def xl_safe(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M UTC")
    if isinstance(v, (list, dict)):
        v = str(v)
    if isinstance(v, str):
        v = v.strip()
        if v[:1] in ("=", "+", "-", "@", "\t", "\r", "|", "%"):
            return "'" + v
    return v


EXPORT_COLUMNS = {
    "contacts": [("Reference", "id"), ("Submitted", "created_at"), ("Name", "name"), ("Email", "email"), ("Phone", "phone"),
                 ("Service", "service"), ("Message", "message"), ("Status", "status"), ("Admin notes", "admin_notes")],
    "assessments": [("Reference", "id"), ("Submitted", "created_at"), ("Completed", "completed_at"), ("Name", "name"), ("Email", "email"),
                    ("Phone", "phone"), ("Status", "status"), ("Recommended services", "recommended"), ("Admin notes", "admin_notes")],
}


@router.get("/export/{kind}")
async def export(kind: str, f: dict = Depends(filters), user: dict = Depends(require_admin)):
    cfg = kind_cfg(kind)
    sort = f.pop("sort")
    q = build_query(cfg, **f)
    docs = await db[cfg["coll"]].find(q, HIDDEN).sort("created_at", 1 if sort == "asc" else -1).to_list(20000)
    wb = Workbook()
    ws = wb.active
    ws.title = kind.capitalize()
    cols = list(EXPORT_COLUMNS[kind])
    max_answers = max((len(d.get("answers") or []) for d in docs), default=0) if kind == "assessments" else 0
    headers = [h for h, _ in cols] + [f"Q{i + 1} / Answer" for i in range(max_answers)]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0E4A8A")
    for d in docs:
        row = [xl_safe(", ".join(d[k]) if k == "recommended" and isinstance(d.get(k), list) else d.get(k, "")) for _, k in cols]
        row += [xl_safe(f"{a.get('question')} — {a.get('answer')}") for a in (d.get("answers") or [])]
        ws.append(row)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    await audit(user, f"export_{kind}", "xlsx", {"rows": len(docs)})
    fname = f"nexcare-{kind}-{now().strftime('%Y%m%d-%H%M')}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"', "Cache-Control": "no-store"})
